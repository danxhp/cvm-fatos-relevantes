#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Recompras e negociações de controlador/administradores (CVM) -> buybacks_cvm.xlsx.

DE ONDE VEM O DADO
------------------
Todo mês cada companhia aberta entrega à CVM, na categoria "Valores Mobiliários
Negociados e Detidos" (art. 11 da Res. CVM 44, antiga ICVM 358), DOIS formulários
em PDF gerados pelo próprio sistema da CVM (layout fixo):

  - "Posição Individual - Cia, Controladas e Coligadas": o que a PRÓPRIA
    companhia (tesouraria) e suas controladas/coligadas negociaram -> recompra.
  - "Posição Consolidada": uma página por grupo — Controlador, Conselho de
    Administração, Diretoria, Conselho Fiscal, Órgãos Técnicos. O grupo vem
    marcado com "( X )" no cabeçalho.

Cada seção traz saldo inicial, movimentações do mês (operação, dia, quantidade,
preço, volume em R$) e saldo final. Prazo: até 10 dias após o fim do mês.

COMO LISTAMOS OS DOCUMENTOS (duas fontes, unidas pelo nº de protocolo)
----------------------------------------------------------------------
  1. Dados abertos da CVM (ipe_cia_aberta_<ano>.zip): histórico completo, mas
     atualizado ~1x por semana — perde o que saiu nos últimos dias.
  2. RAD (o mesmo endpoint do monitor), consultando os últimos
     RAD_LOOKBACK_DAYS dias por intervalo de datas: cobre essa defasagem.

REAPRESENTAÇÕES
---------------
Empresas reapresentam meses antigos (ex.: HAPV3 reentregou 07/2026 em outubro).
Por (ticker, mês de referência, formulário) vale SÓ o documento mais recente; os
anteriores aparecem na aba Documentos como "substituído".

O QUE ENTRA NAS ABAS DE RESUMO
------------------------------
Só AÇÕES negociadas no mercado: compra/venda à vista ou a termo. Transferência
por vesting/ILP, liquidação de swap, empréstimo de ações, debêntures, opções etc.
ficam só na aba Detalhe (Operação = "Outros"). Administradores = Conselho de
Administração + Diretoria (definição da Lei 6.404); Conselho Fiscal e Órgãos
Técnicos vão para "Outros".

CONFERÊNCIA EMBUTIDA
--------------------
Para cada seção: saldo inicial + compras - vendas tem que bater com o saldo
final (em ações). Diferença aparece na aba Documentos. Ex.: controlador da HAPV3
em 08/2026: 196.263.730 + 910.500 compradas + 4 x 89.048 recebidas por swap =
197.530.422, exatamente o saldo final declarado.

Uso:
    py buybacks_cvm.py               # atualiza: lista, baixa só o novo, regera o xlsx
    py buybacks_cvm.py --reparse     # relê todos os PDFs (depois de mudar o parser)
    py buybacks_cvm.py --dump-ops    # lista as descrições de operação distintas e a
                                     # classificação de cada uma (auditar regras)
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import time
import unicodedata
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from urllib.parse import parse_qs, urlencode, urlparse

import pypdf
import requests

import cvm_fatos_relevantes_claude as cvm


SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT_XLSX = SCRIPT_DIR / "buybacks_cvm.xlsx"
# Resultado do parse por protocolo. Versionado: a próxima execução (em qualquer
# máquina) só baixa e lê o que ainda não está aqui.
CACHE_FILE = SCRIPT_DIR / "buybacks_cache.json"
# PDFs baixados. NÃO versionado (.gitignore) — só serve para reparse local.
PDF_DIR = SCRIPT_DIR / "buybacks_pdfs"

START_REF = "2024-10"      # primeiro mês de referência (AAAA-MM): 24 meses = LTM atual + anterior
RAD_LOOKBACK_DAYS = 45     # janela do RAD; cobre com folga a defasagem dos dados abertos
# Suba quando mudar a lógica de parse: entradas do cache com versão antiga são relidas.
PARSER_VERSION = 2  # 2: "Qualificação" com o valor na linha seguinte (layout de 09/2026)

IPE_URL = "https://dados.cvm.gov.br/dados/CIA_ABERTA/DOC/IPE/DADOS/ipe_cia_aberta_{year}.zip"
CATEGORY_PREFIX = "valores mobiliarios negociados"  # comparado sem acento/caixa

FORM_INDIVIDUAL = "Individual"
FORM_CONSOLIDADO = "Consolidado"

# Ordem das caixas "( )" no cabeçalho do formulário consolidado.
GRUPOS_CONSOLIDADO = [
    "Controlador",
    "Conselho de Administração",
    "Diretoria",
    "Conselho Fiscal",
    "Órgãos Técnicos ou Consultivos",
]

ABA_EMPRESA = "Empresa"
ABA_CONTROLADORES = "Controladores"
ABA_ADMINISTRADORES = "Administradores"
ABA_OUTROS = "Outros"

OP_COMPRA = "Compra"
OP_VENDA = "Venda"
OP_OUTROS = "Outros"
# Compra/venda declarada sem preço: a quantidade é real (o saldo confirma), o
# valor não. Fica fora dos SUMIFS da planilha (que filtram "Compra"/"Venda").
OP_COMPRA_SP = "Compra s/ preço"
OP_VENDA_SP = "Venda s/ preço"
# Par compra a termo + venda à vista no mesmo dia e quantidade: posição igual.
OP_ROLAGEM = "Rolagem a termo"

# Qualquer uma destas palavras na descrição faz a operação NÃO ser compra/venda de
# mercado, mesmo que contenha "compra"/"venda" (ex.: "exercício de opção de compra",
# "liquidação de venda a termo"). Comparadas sem acento, em minúsculas.
NAO_MERCADO = (
    "vesting", "plano", "ilp", "pilp", "swap", "emprestimo", "aluguel", "bonifica",
    "subscri", "exerc", "opcao", "doacao", "heranca", "transfer", "recebimento",
    "entrega", "desdobr", "grupament", "cancelamento", "incorpora", "conversao",
    "permuta", "dacao", "liquidacao", "restrit", "matching", "outorga", "remunera",
    "posse", "ajuste", "devolucao", "contratacao", "aliena", "fraco", "leilao",
    "contribui", "judicial", "penhora", "homologa", "integraliza",
    # mudança de cargo: a posição entra/sai do grupo, ninguém negociou
    "desligamento", "saida", "entrada", "eleicao", "renuncia",
)


def log(msg: str) -> None:
    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}")


def fold(s: str) -> str:
    """minúsculas e sem acento — para comparar texto livre do formulário."""
    s = unicodedata.normalize("NFKD", s or "")
    return "".join(c for c in s if not unicodedata.combining(c)).lower()


def br_int(s: str) -> int:
    return int(s.replace(".", ""))


def br_float(s: str) -> float:
    return float(s.replace(".", "").replace(",", "."))


# ============================================================================
# LISTAGEM DOS DOCUMENTOS
# ============================================================================

def _form_from_tipo(tipo: str) -> Optional[str]:
    t = fold(tipo)
    if "consolidada" in t:
        return FORM_CONSOLIDADO
    if "individual" in t:
        return FORM_INDIVIDUAL
    return None


def list_from_open_data(session: requests.Session, years: List[int]) -> Dict[str, Dict]:
    docs: Dict[str, Dict] = {}
    for year in years:
        try:
            r = session.get(IPE_URL.format(year=year), timeout=120)
            r.raise_for_status()
            # Página de manutenção com HTTP 200 também cai aqui (BadZipFile),
            # em vez de derrubar a execução inteira.
            z = zipfile.ZipFile(io.BytesIO(r.content))
            with z.open(z.namelist()[0]) as fh:
                rows = list(csv.DictReader(io.TextIOWrapper(fh, encoding="latin-1"), delimiter=";"))
        except Exception as e:
            log(f"WARNING: dados abertos {year} indisponíveis: {e}")
            continue
        for row in rows:
            if not fold(row.get("Categoria", "")).startswith(CATEGORY_PREFIX):
                continue
            company = cvm.COMPANY_CODES.get(cvm.normalize_cvm_code(row.get("Codigo_CVM")))
            ref = (row.get("Data_Referencia") or "")[:7]
            form = _form_from_tipo(row.get("Tipo", ""))
            if not company or not form or ref < START_REF:
                continue
            q = parse_qs(urlparse(row.get("Link_Download", "")).query)
            prot = (q.get("numProtocolo") or [""])[0]
            if not prot:
                continue
            docs[prot] = {
                "protocolo": prot,
                "sequencia": (q.get("numSequencia") or [""])[0],
                "versao": (q.get("numVersao") or [row.get("Versao", "")])[0],
                "ticker": company["ticker"],
                "empresa": company["name"],
                "formulario": form,
                "ref": ref,
                "entrega": row.get("Data_Entrega", ""),
                "apresentacao": (row.get("Tipo_Apresentacao") or "")[:2],
            }
    return docs


def list_from_rad(session: requests.Session, days: int) -> Dict[str, Dict]:
    """Mesmo endpoint do monitor, mas por intervalo de datas (periodo=2)."""
    hoje = date.today()
    payload = {
        "dataDe": (hoje - timedelta(days=days)).strftime("%d/%m/%Y"),
        "dataAte": hoje.strftime("%d/%m/%Y"),
        "empresa": "", "setorAtividade": "-1", "categoriaEmissor": "-1",
        "situacaoEmissor": "A", "tipoParticipante": "1", "dataReferencia": "",
        "categoria": cvm.CVM_ALL_CATEGORIES, "periodo": "2", "horaIni": "", "horaFim": "",
        "palavraChave": "", "ultimaDtRef": "false", "tipoEmpresa": "2",
        "token": "", "versaoCaptcha": "V3",
    }
    resp = session.post(
        cvm.LIST_URL, json=payload,
        headers={"Content-Type": "application/json; charset=utf-8",
                 "X-Requested-With": "XMLHttpRequest"},
        timeout=120,
    )
    resp.raise_for_status()
    d = resp.json().get("d", {})
    if isinstance(d, dict) and d.get("temErro"):
        raise RuntimeError(f"RAD/CVM: {d.get('msgErro')}")
    raw = d.get("dados", "") if isinstance(d, dict) else (d or "")

    docs: Dict[str, Dict] = {}
    for row in raw.split("&*"):
        cols = row.split("$&")
        if len(cols) < 11:
            continue
        if not fold(cvm.strip_html(cols[2])).startswith(CATEGORY_PREFIX):
            continue
        company = cvm.COMPANY_CODES.get(cvm.normalize_cvm_code(cvm.strip_html(cols[0])))
        form = _form_from_tipo(cvm.strip_html(cols[3]))
        # col 5 = "20260901 09/2026"; col 6 = "20261009 09/10/2026 16:44"
        mref = re.match(r"(\d{4})(\d{2})", cvm.strip_html(cols[5]))
        ment = re.search(r"(\d{2})/(\d{2})/(\d{4})\s+(\d{2}:\d{2})", cvm.strip_html(cols[6]))
        mdl = re.search(
            r"OpenDownloadDocumentos\(\s*'(\d+)'\s*,\s*'(\d+)'\s*,\s*'(\d+)'", cols[10])
        if not (company and form and mref and mdl):
            continue
        ref = f"{mref.group(1)}-{mref.group(2)}"
        if ref < START_REF:
            continue
        seq, ver, prot = mdl.groups()
        docs[prot] = {
            "protocolo": prot,
            "sequencia": seq,
            "versao": ver,
            "ticker": company["ticker"],
            "empresa": company["name"],
            "formulario": form,
            "ref": ref,
            "entrega": (f"{ment.group(3)}-{ment.group(2)}-{ment.group(1)} {ment.group(4)}"
                        if ment else ""),
            "apresentacao": cvm.strip_html(cols[9])[:2] if len(cols) > 9 else "",
        }
    return docs


def list_documents(session: requests.Session) -> Dict[str, Dict]:
    years = list(range(int(START_REF[:4]), date.today().year + 1))
    docs = list_from_open_data(session, years)
    n_open = len(docs)
    try:
        rad = list_from_rad(session, RAD_LOOKBACK_DAYS)
    except Exception as e:
        # Sem o RAD perdemos só a última semana; o resto vem dos dados abertos.
        log(f"WARNING: RAD indisponível ({e}) — documentos dos últimos dias podem faltar.")
        rad = {}
    for prot, d in rad.items():
        if prot in docs:
            docs[prot]["entrega"] = d["entrega"] or docs[prot]["entrega"]  # RAD tem hora
        else:
            docs[prot] = d
    log(f"Documentos: {n_open} nos dados abertos, {len(rad)} no RAD "
        f"({RAD_LOOKBACK_DAYS}d), {len(docs)} no total.")
    return docs


def download_url(doc: Dict) -> str:
    return cvm.DOWNLOAD_URL + "?" + urlencode({
        "Tela": "ext", "numSequencia": doc["sequencia"], "numVersao": doc["versao"],
        "numProtocolo": doc["protocolo"], "descTipo": "IPE", "CodigoInstituicao": "1",
    })


def viewer_url(doc: Dict) -> str:
    return cvm.VIEWER_URL + "?" + urlencode({"NumeroProtocoloEntrega": doc["protocolo"]})


def fetch_pdf(session: requests.Session, doc: Dict) -> bytes:
    PDF_DIR.mkdir(exist_ok=True)
    path = PDF_DIR / f"{doc['ticker']}_{doc['ref']}_{doc['protocolo']}.pdf"
    if path.exists():
        return path.read_bytes()
    erro: Optional[Exception] = None
    for tentativa in range(3):  # o RAD oscila; 3 tentativas antes de desistir
        try:
            r = session.get(download_url(doc), timeout=90)
            r.raise_for_status()
            if r.content[:4] != b"%PDF":
                raise RuntimeError("resposta não é PDF")
            path.write_bytes(r.content)
            return r.content
        except Exception as e:
            erro = e
            time.sleep(2 * (tentativa + 1))
    raise RuntimeError(f"download falhou após 3 tentativas: {erro}")


# ============================================================================
# PARSE DO PDF
# ============================================================================

_INT = r"-?\d{1,3}(?:\.\d{3})*"
_DEC = r"-?\d{1,3}(?:\.\d{3})*,\d+"
# Uma linha de movimentação termina SEMPRE em "dia quantidade preço volume". O
# texto que vem antes (valor mobiliário, características, intermediário,
# operação) quebra linha de forma diferente em cada PDF, então não dá para
# confiar em posição — ancoramos no rabo numérico.
_ROW_RE = re.compile(
    rf"(?P<desc>.*?)\s*(?<!\S)(?P<dia>[0-3]?\d)\s+(?P<qtd>{_INT})\s+"
    rf"(?P<preco>{_DEC})\s+(?P<vol>{_DEC})(?!\S)"
)
_SALDO_RE = re.compile(rf"(?P<desc>.*?)\s*(?<!\S)(?P<qtd>{_INT})(?!\S)")
_SALDO_HEADER_RE = re.compile(
    r"Valor\s+Mobili[áa]rio\s*(?:/\s*)?Derivativo\s+Caracter[íi]sticas\s+dos\s+"
    r"T[íi]tulos\s+Quantidade", re.I)
# Onde começa uma linha da tabela: o valor mobiliário + característica. Precisa
# disso porque, quando uma linha quebra na virada de página, o pypdf devolve o
# resto da descrição DEPOIS dos números ("...plano de 30 3.280 34,52 113.225,60
# remuneração Ações ON Ágora ...") — sem cortar aqui, "remuneração" iria parar
# na descrição da linha seguinte e poderia mudar a classificação dela.
_ACAO_START_RE = re.compile(r"(?i:a[çc][õo]es|units?)\s+(?:ON|PN[A-Z]?|UNT|UNIT)\b")
_OUTRO_START_RE = re.compile(
    r"Deb[êe]ntures\b|Derivativos?\b|Op[çc][ãa]o\s+de\s+(?:Compra|Venda)\b|"
    r"B[ôo]nus\s+de\s+Subscri|Direitos?\s+de\s+Subscri|Recibos?\s+de\s+Subscri|"
    r"BDRs?\b|Notas?\s+Promiss|Cotas?\b|Outros\s")  # "Outros" = 1a coluna (ex.: ADR)
# Sobra com cara de linha numérica não lida = layout inesperado (vira alerta).
_NUMERIC_LEFTOVER_RE = re.compile(r"\d+,\d+")


def _squash(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def _parse_saldo(text: str) -> List[Tuple[str, int]]:
    text = _squash(_SALDO_HEADER_RE.sub(" ", text))
    return [(_squash(m.group("desc")), br_int(m.group("qtd")))
            for m in _SALDO_RE.finditer(text) if m.group("desc").strip()]


def _row_start(desc: str) -> int:
    starts = [m.start() for m in (_ACAO_START_RE.search(desc), _OUTRO_START_RE.search(desc)) if m]
    return min(starts) if starts else 0


def _parse_movimentos(text: str) -> Tuple[List[Dict], str]:
    """Devolve (operações, sobra). Sobra != '' = texto com números que o regex
    não consumiu (layout inesperado) — vira alerta na aba Documentos."""
    text = _squash(text)
    ops: List[Dict] = []
    pos = 0
    for m in _ROW_RE.finditer(text):
        desc = m.group("desc")
        cut = _row_start(desc)
        if cut and ops:  # o começo é o fim da descrição da linha anterior
            ops[-1]["desc"] = _squash(ops[-1]["desc"] + " " + desc[:cut])
            desc = desc[cut:]
        ops.append({
            "desc": _squash(desc),
            "dia": int(m.group("dia")),
            "qtd": br_int(m.group("qtd")),
            "preco": br_float(m.group("preco")),
            "volume": br_float(m.group("vol")),
        })
        pos = m.end()
    sobra = text[pos:].strip()
    if ops and sobra and not _NUMERIC_LEFTOVER_RE.search(sobra):
        ops[-1]["desc"] = _squash(ops[-1]["desc"] + " " + sobra)
        sobra = ""
    return ops, sobra


def _grupo_consolidado(header: str) -> Optional[str]:
    boxes = re.findall(r"\(\s*([Xx]?)\s*\)", header)
    marked = [i for i, b in enumerate(boxes) if b]
    if len(boxes) == len(GRUPOS_CONSOLIDADO) and len(marked) == 1:
        return GRUPOS_CONSOLIDADO[marked[0]]
    return None


def parse_pdf(data: bytes) -> Dict:
    reader = pypdf.PdfReader(io.BytesIO(data))
    text = "\n".join((p.extract_text() or "") for p in reader.pages)

    head = fold(text[:400])
    formulario = (FORM_CONSOLIDADO if "consolidado" in head
                  else FORM_INDIVIDUAL if "individual" in head else None)
    mref = re.search(r"\bEm\s+(\d{2})/(\d{4})", text)
    out = {
        "parser": PARSER_VERSION,
        "formulario_pdf": formulario,
        "ref_pdf": f"{mref.group(2)}-{mref.group(1)}" if mref else None,
        "secoes": [],
    }

    # Uma seção por "Denominação da ...": no consolidado, uma por grupo; no
    # individual, uma por entidade (companhia/tesouraria, cada controlada...).
    # ATENÇÃO: o consolidado também traz seções "Denominação da Controladora" —
    # posição dos administradores em ações da CONTROLADORA (ex.: conselheiro da
    # VIVT3 com ações da Telefónica S.A.). Não é ação da companhia: "emissor"
    # guarda isso e essas seções nunca entram nos resumos.
    parts = re.split(
        r"Denomina[çc][ãa]o da (Companhia|Controladora|Controlada|Coligada)\s*:?", text)
    for emissor, chunk in zip(parts[1::2], parts[2::2]):
        # o cabeçalho da próxima página pode vazar no fim do pedaço
        chunk = re.split(r"FORMUL[ÁA]RIO\s+(?:CONSOLIDADO|INDIVIDUAL)", chunk)[0]
        if "Saldo Inicial" not in chunk or "Saldo Final" not in chunk:
            continue
        header, rest = chunk.split("Saldo Inicial", 1)
        saldo_ini_txt, rest = (rest.split("Movimentações no Mês", 1)
                               if "Movimentações no Mês" in rest else ("", rest))
        mov_txt, saldo_fim_txt = rest.split("Saldo Final", 1)
        mov_txt = mov_txt.split("Volume (R$)", 1)[1] if "Volume (R$)" in mov_txt else ""

        sec: Dict = {
            "emissor": emissor,
            "emissor_nome": _squash(header.strip().split("\n", 1)[0]),
            "grupo": None,
            "nome": None,
        }
        if formulario == FORM_CONSOLIDADO:
            g = header.split("Grupo e Pessoas", 1)
            sec["grupo"] = _grupo_consolidado(g[1] if len(g) > 1 else header)
        else:
            # \s* atravessa a quebra: no layout de 09/2026 o valor vem na linha de
            # baixo ("Qualificação:\nColigada"). Em branco continua dando "", porque
            # `header` termina antes de "Saldo Inicial".
            mq = re.search(r"Qualifica[çc][ãa]o:\s*([^\n]*)", header)
            mn = re.search(r"Nome:\s*(.+?)\s*CPF/CNPJ", header, re.S)
            # Várias companhias deixam "Qualificação:" em branco; vazio != erro.
            sec["grupo"] = (_squash(mq.group(1)) or "(qualificação em branco)") if mq else None
            sec["nome"] = _squash(mn.group(1)) if mn else None

        ops, sobra = _parse_movimentos(mov_txt)
        sec["saldo_inicial"] = _parse_saldo(saldo_ini_txt)
        sec["saldo_final"] = _parse_saldo(saldo_fim_txt)
        sec["ops"] = ops
        sec["sobra"] = sobra
        out["secoes"].append(sec)
    return out


# ============================================================================
# CLASSIFICAÇÃO
# ============================================================================

def valor_mobiliario(desc: str) -> str:
    f = fold(desc)
    if f.startswith(("acoes", "acao", "units", "unit ")):
        return "Ações"
    if f.startswith("debentures"):
        return "Debêntures"
    if f.startswith("derivativo"):
        return "Derivativos"
    if f.startswith("opcao"):
        return "Opções"
    return desc.split(" ")[0] if desc else ""


def motivo_nao_mercado(desc: str) -> Optional[str]:
    """Palavra de NAO_MERCADO presente na descrição (None se não houver)."""
    f = fold(desc)
    return next((k for k in NAO_MERCADO if k in f), None)


def classificar_operacao(op: Dict) -> str:
    f = fold(op["desc"])
    if motivo_nao_mercado(op["desc"]):
        return OP_OUTROS
    compra = bool(re.search(r"\bcompras?\b", f))
    venda = bool(re.search(r"\bvendas?\b", f))
    # A YDUQ3 (10/2024 e 12/2024) escreveu só "RECOMPRA DE ACOES". Testado depois
    # de "venda" para "venda ... programa de recompra" continuar sendo Venda.
    if not compra and not venda and re.search(r"\brecompras?\b", f):
        compra = True
    if not (compra or venda):
        return OP_OUTROS
    # Compra/venda declarada SEM preço (ONCO3 12/2025, ODPV3: "Venda à vista ...
    # 0,00000 0,00"): o saldo mostra que as ações saíram, então a quantidade
    # conta, mas o valor não é conhecido — classe própria, fora do preço médio.
    sem_preco = abs(op["preco"]) == 0 or abs(op["volume"]) == 0
    if compra:
        return OP_COMPRA_SP if sem_preco else OP_COMPRA
    return OP_VENDA_SP if sem_preco else OP_VENDA


def rolagens_a_termo(ops: List[Dict]) -> set:
    """Índices das linhas de "compra a termo" + "venda à vista" (ou o inverso)
    do MESMO dia, na mesma seção, quando as duas pontas somam a mesma
    quantidade. A posição não muda: é captação/rolagem via termo (HYPE3:
    diretores vendem à vista e recompram a termo — às vezes em várias linhas,
    4.392 + 40.608 a termo contra 45.000 à vista), não compra nem venda."""
    pontas: Dict[Tuple[int, str], List[int]] = {}
    for i, op in enumerate(ops):
        if valor_mobiliario(op["desc"]) != "Ações":
            continue
        cls = classificar_operacao(op)
        if cls not in (OP_COMPRA, OP_VENDA):
            continue
        termo = "termo" in fold(op["desc"])
        pontas.setdefault((op["dia"], f"{cls}-{'termo' if termo else 'vista'}"), []).append(i)
    idx = set()
    for (dia, chave), linhas in pontas.items():
        par = {"Compra-termo": "Venda-vista", "Venda-termo": "Compra-vista"}.get(chave)
        outras = pontas.get((dia, par)) if par else None
        if not outras:
            continue
        soma = lambda ls: sum(abs(ops[i]["qtd"]) for i in ls)
        if soma(linhas) == soma(outras) > 0:
            idx.update(linhas)
            idx.update(outras)
    return idx


def aba_da_secao(formulario: str, sec: Dict) -> str:
    if sec.get("emissor") != "Companhia":
        return ABA_OUTROS  # ações da controladora/controlada, não da companhia
    g = fold(sec.get("grupo") or "")
    if formulario == FORM_INDIVIDUAL:
        # Companhia (tesouraria) e controladas = recompra consolidada. Coligada
        # não é controlada pela companhia, então não é recompra.
        return ABA_OUTROS if "coligada" in g else ABA_EMPRESA
    if g == fold("Controlador"):
        return ABA_CONTROLADORES
    if g in (fold("Conselho de Administração"), fold("Diretoria")):
        return ABA_ADMINISTRADORES
    return ABA_OUTROS


def descrever_grupo(formulario: str, sec: Dict) -> str:
    g = sec.get("grupo") or "?"
    if formulario == FORM_INDIVIDUAL and sec.get("nome"):
        g = f"{g} — {sec['nome']}"
    if sec.get("emissor") != "Companhia":
        g = f"{g} — ações da {fold(sec.get('emissor') or '')} {sec.get('emissor_nome') or ''}"
    return g.strip()


def conferir_secao(sec: Dict) -> Optional[Tuple[int, int]]:
    """(diferença, qtd não-mercado) quando saldo inicial + compras - vendas
    não bate com o saldo final, em ações; None quando bate."""
    def acoes(saldo):
        return sum(q for d, q in saldo if valor_mobiliario(d) == "Ações")
    liq = outros = 0
    for op in sec["ops"]:
        if valor_mobiliario(op["desc"]) != "Ações":
            continue
        cls = classificar_operacao(op)
        if cls in (OP_COMPRA, OP_COMPRA_SP):
            liq += abs(op["qtd"])
        elif cls in (OP_VENDA, OP_VENDA_SP):
            liq -= abs(op["qtd"])
        elif motivo_nao_mercado(op["desc"]) or not op["preco"] or not op["volume"]:
            # há quem informe a saída com sinal (HAPV3: entrega de ações de plano
            # como "-217.984" a "-12,72"); o que importa aqui é o tamanho.
            outros += abs(op["qtd"])
        # Linha com preço e volume, sem rótulo de compra/venda e sem motivo
        # não-mercado NÃO cobre diferença: foi assim que "RECOMPRA DE ACOES" da
        # YDUQ3 passou como OK*. Sem cobertura, o documento vira VERIFICAR.
    dif = acoes(sec["saldo_final"]) - acoes(sec["saldo_inicial"]) - liq
    return (dif, outros) if dif else None


# ============================================================================
# PLANILHA
# ============================================================================

def _fmt_int(n: int) -> str:
    return f"{n:+,}".replace(",", ".")


def vigentes(cache: Dict[str, Dict]) -> Dict[Tuple[str, str, str], str]:
    """(ticker, mês, formulário) -> protocolo vigente. Reapresentação: vale o
    protocolo mais recente (o número de protocolo da CVM é sequencial no tempo)."""
    latest: Dict[Tuple[str, str, str], str] = {}
    for p, e in cache.items():
        m = e["meta"]
        if m["ref"] < START_REF:
            continue
        k = (m["ticker"], m["ref"], m["formulario"])
        if k not in latest or int(p) > int(latest[k]):
            latest[k] = p
    return latest


def documento_lido(entry: Dict) -> bool:
    """O PDF foi baixado e lido? Download que falhou (entry.parse.erro) ou PDF
    sem nenhuma seção da companhia NÃO pode virar "zero recompra"."""
    parsed = entry.get("parse") or {}
    return not parsed.get("erro") and any(
        s.get("emissor") == "Companhia" for s in parsed.get("secoes", []))


def saldo_tesouraria(entry: Dict) -> Optional[int]:
    """Ações em tesouraria no fim do mês (companhia + controladas), do saldo final
    do formulário individual. Coligada fica de fora, como nas regras de recompra.
    None se o PDF não foi lido."""
    secoes = [s for s in entry["parse"].get("secoes", [])
              if aba_da_secao(FORM_INDIVIDUAL, s) == ABA_EMPRESA]
    if not documento_lido(entry) or not secoes:
        return None
    return sum(q for s in secoes for d, q in s["saldo_final"] if valor_mobiliario(d) == "Ações")


def build_rows(cache: Dict[str, Dict]) -> Tuple[List[Dict], List[Dict]]:
    """Devolve (operações para a aba Detalhe, linhas para a aba Documentos)."""
    docs = {p: e for p, e in cache.items() if e["meta"]["ref"] >= START_REF}
    latest = vigentes(cache)

    ops_rows: List[Dict] = []
    doc_rows: List[Dict] = []
    for p, e in docs.items():
        m, parsed = e["meta"], e["parse"]
        k = (m["ticker"], m["ref"], m["formulario"])
        usado = latest[k] == p
        secoes = parsed.get("secoes", [])

        alertas: List[str] = []
        if parsed.get("erro"):
            alertas.append(f"ERRO: {parsed['erro']}")
        elif not secoes:
            alertas.append("NÃO LIDO: nenhuma seção reconhecida no PDF — conferir manualmente")
        if parsed.get("ref_pdf") and parsed["ref_pdf"] != m["ref"]:
            alertas.append(f"mês no PDF ({parsed['ref_pdf']}) difere do cadastro")
        so_explicadas = True
        for sec in secoes:
            if sec.get("sobra"):
                alertas.append(f"texto não lido: {sec['sobra'][:80]}")
                so_explicadas = False
            if sec.get("emissor") != "Companhia":
                continue
            r = conferir_secao(sec)
            if r:
                dif, outros = r
                explicada = outros and abs(dif) <= outros
                so_explicadas = so_explicadas and bool(explicada)
                alertas.append(
                    f"{sec.get('grupo')}: saldo final {_fmt_int(dif)} ações vs. compras/vendas"
                    + (" (coberto por operações não-mercado)" if explicada else ""))
        if not alertas:
            conferencia = "OK"
        elif so_explicadas and secoes and not parsed.get("erro"):
            conferencia = "OK* " + "; ".join(alertas)
        else:
            conferencia = "VERIFICAR: " + "; ".join(alertas)

        n_ops = sum(len(s["ops"]) for s in secoes)
        doc_rows.append({
            "meta": m, "usado": usado, "substituto": None if usado else latest[k],
            "n_ops": n_ops, "conferencia": conferencia,
        })
        if not usado:
            continue
        for sec in secoes:
            aba = aba_da_secao(m["formulario"], sec)
            grupo = descrever_grupo(m["formulario"], sec)
            rolagem = rolagens_a_termo(sec["ops"])
            for j, op in enumerate(sec["ops"]):
                cls = OP_ROLAGEM if j in rolagem else classificar_operacao(op)
                row = {"meta": m, "aba": aba, "grupo": grupo,
                       "vm": valor_mobiliario(op["desc"]), "op": cls, **op}
                if cls in (OP_COMPRA, OP_VENDA, OP_COMPRA_SP, OP_VENDA_SP, OP_ROLAGEM):
                    # A direção está na coluna Operação; um sinal informado no
                    # formulário faria o SUMIFS de "Venda" sair negativo.
                    for f in ("qtd", "preco", "volume"):
                        row[f] = abs(row[f])
                ops_rows.append(row)

    # Cobertura: (ticker, mês, formulário) esperado e não encontrado. Sem isso,
    # "não aparece na aba Empresa" seria ambíguo entre "não recomprou" e "ainda
    # não entregou".
    refs = sorted({e["meta"]["ref"] for e in docs.values()})
    tickers = {c["ticker"]: c["name"] for c in cvm.COMPANIES if c.get("cvm_code")}
    for ref in refs:
        for t, name in tickers.items():
            for form in (FORM_INDIVIDUAL, FORM_CONSOLIDADO):
                if (t, ref, form) not in latest:
                    doc_rows.append({
                        "meta": {"ticker": t, "empresa": name, "ref": ref, "formulario": form,
                                 "protocolo": "", "versao": "", "apresentacao": "", "entrega": ""},
                        "usado": False, "substituto": None, "n_ops": None,
                        "conferencia": "NÃO ENCONTRADO (prazo: 10 dias após o fim do mês)",
                    })
    return ops_rows, doc_rows


def write_xlsx(ops_rows: List[Dict], doc_rows: List[Dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    font = Font(name="Arial", size=10)
    font_head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_link = Font(name="Arial", size=10, color="0563C1", underline="single")
    fill_head = PatternFill("solid", fgColor="1F3864")
    fill_warn = PatternFill("solid", fgColor="FFF2CC")
    align_head = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def ref_date(ref: str) -> datetime:
        return datetime(int(ref[:4]), int(ref[5:7]), 1)

    def header(ws, cols: List[Tuple[str, int]]) -> None:
        for i, (title, width) in enumerate(cols, 1):
            c = ws.cell(row=1, column=i, value=title)
            c.font, c.fill, c.alignment = font_head, fill_head, align_head
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

    def finish(ws, n_cols: int) -> None:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.hyperlink is None:
                    c.font = font
        last = max(ws.max_row, 2)
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{last}"

    wb = Workbook()
    wb.remove(wb.active)
    # Fórmulas sem valor em cache (openpyxl não calcula): força o Excel a
    # recalcular tudo ao abrir.
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    # --- Abas de resumo: fórmulas SUMIFS sobre a aba Detalhe -------------------
    nota = ("Soma das operações da aba Detalhe com Aba = '{aba}', Valor mobiliário = "
            "'Ações' e Operação = Compra ou Venda (negociação de mercado à vista/termo). "
            "Vesting, ILP, swap, empréstimo, subscrição etc. ficam fora (Operação = Outros).")
    resumo_cols = [("Empresa", 26), ("Ticker", 9), ("Mês e ano", 11), ("Operação", 10),
                   ("Qtd movimentada", 16), ("Preço médio (R$)", 15), ("Valor total (R$)", 18)]
    for aba in (ABA_EMPRESA, ABA_CONTROLADORES, ABA_ADMINISTRADORES):
        ws = wb.create_sheet(aba)
        header(ws, resumo_cols)
        ws["E1"].comment = Comment(nota.format(aba=aba), "buybacks_cvm.py")
        ws["F1"].comment = Comment("Preço médio ponderado = Valor total / Qtd movimentada.",
                                   "buybacks_cvm.py")
        combos = sorted(
            {(r["meta"]["ref"], r["meta"]["ticker"], r["meta"]["empresa"], r["op"])
             for r in ops_rows
             if r["aba"] == aba and r["vm"] == "Ações" and r["op"] in (OP_COMPRA, OP_VENDA)},
            key=lambda x: (-int(x[0].replace("-", "")), x[2], x[3]),
        )
        crit = (f'Detalhe!$D:$D,"{aba}",Detalhe!$B:$B,$B{{r}},Detalhe!$C:$C,$C{{r}},'
                f'Detalhe!$H:$H,$D{{r}},Detalhe!$F:$F,"Ações"')
        for i, (ref, ticker, empresa, op) in enumerate(combos, 2):
            ws.cell(row=i, column=1, value=empresa)
            ws.cell(row=i, column=2, value=ticker)
            ws.cell(row=i, column=3, value=ref_date(ref)).number_format = "mm/yyyy"
            ws.cell(row=i, column=4, value=op)
            ws.cell(row=i, column=5, value=f"=SUMIFS(Detalhe!$J:$J,{crit.format(r=i)})"
                    ).number_format = "#,##0"
            ws.cell(row=i, column=6, value=f"=IF(E{i}=0,0,G{i}/E{i})").number_format = "#,##0.00"
            ws.cell(row=i, column=7, value=f"=SUMIFS(Detalhe!$L:$L,{crit.format(r=i)})"
                    ).number_format = "#,##0"
        finish(ws, len(resumo_cols))

    # --- Detalhe: uma linha por operação lida (só documentos vigentes) ---------
    ws = wb.create_sheet("Detalhe")
    det_cols = [("Empresa", 22), ("Ticker", 9), ("Mês e ano", 11), ("Aba", 15),
                ("Grupo / qualificação", 40), ("Valor mobiliário", 13),
                ("Descrição no formulário", 60), ("Operação", 10), ("Dia", 6),
                ("Quantidade", 14), ("Preço (R$)", 13), ("Volume (R$)", 16),
                ("Protocolo", 11), ("Documento", 11)]
    header(ws, det_cols)
    ops_sorted = sorted(ops_rows, key=lambda r: (
        -int(r["meta"]["ref"].replace("-", "")), r["meta"]["empresa"], r["aba"], r["grupo"],
        r["dia"]))
    for i, r in enumerate(ops_sorted, 2):
        m = r["meta"]
        vals = [m["empresa"], m["ticker"], ref_date(m["ref"]), r["aba"], r["grupo"], r["vm"],
                r["desc"], r["op"], r["dia"], r["qtd"], r["preco"], r["volume"],
                int(m["protocolo"])]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=3).number_format = "mm/yyyy"
        ws.cell(row=i, column=10).number_format = "#,##0"
        ws.cell(row=i, column=11).number_format = "#,##0.00###"
        ws.cell(row=i, column=12).number_format = "#,##0.00"
        ws.cell(row=i, column=13).number_format = "0"
        link = ws.cell(row=i, column=14, value="abrir")
        link.hyperlink, link.font = viewer_url(m), font_link
    finish(ws, len(det_cols))

    # --- Documentos: tudo que foi listado, vigente ou não, e o que falta -------
    ws = wb.create_sheet("Documentos")
    doc_cols = [("Empresa", 22), ("Ticker", 9), ("Mês e ano", 11), ("Formulário", 12),
                ("Protocolo", 11), ("Versão", 8), ("Apresentação", 12), ("Entrega", 17),
                ("Situação", 26), ("Operações lidas", 10), ("Conferência de saldo", 70),
                ("Documento", 11)]
    header(ws, doc_cols)
    docs_sorted = sorted(doc_rows, key=lambda d: (
        -int(d["meta"]["ref"].replace("-", "")), d["meta"]["empresa"], d["meta"]["formulario"],
        d["meta"]["protocolo"]))
    for i, d in enumerate(docs_sorted, 2):
        m = d["meta"]
        situacao = ("" if not m["protocolo"] else "Vigente" if d["usado"]
                    else f"Substituído pelo {d['substituto']}")
        vals = [m["empresa"], m["ticker"], ref_date(m["ref"]), m["formulario"],
                int(m["protocolo"]) if m["protocolo"] else None, m["versao"],
                {"AP": "Apresentação", "RE": "Reapresentação"}.get(m["apresentacao"],
                                                                  m["apresentacao"]),
                m["entrega"], situacao, d["n_ops"], d["conferencia"]]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=3).number_format = "mm/yyyy"
        ws.cell(row=i, column=5).number_format = "0"
        if m["protocolo"]:
            link = ws.cell(row=i, column=12, value="abrir")
            link.hyperlink, link.font = viewer_url(m), font_link
        if (d["usado"] and not d["conferencia"].startswith("OK")) or not m["protocolo"]:
            for j in range(1, len(doc_cols) + 1):
                ws.cell(row=i, column=j).fill = fill_warn
    finish(ws, len(doc_cols))

    # --- Leia-me ---------------------------------------------------------------
    ws = wb.create_sheet("Leia-me")
    ws.column_dimensions["A"].width = 120
    linhas = [
        "Recompras (empresa) e negociações de controladores e administradores — CVM",
        f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} por buybacks_cvm.py "
        f"(meses de referência a partir de {START_REF[5:]}/{START_REF[:4]}).",
        "",
        "FONTE: formulários mensais 'Valores Mobiliários Negociados e Detidos' (art. 11 da "
        "Res. CVM 44) entregues à CVM. Lista de documentos: dados abertos da CVM "
        "(dados.cvm.gov.br, IPE) + consulta ao RAD dos últimos dias.",
        "",
        "ABAS",
        "Empresa — formulário 'Posição Individual': a própria companhia (tesouraria) e "
        "controladas. Compra = recompra; venda = alienação de ações em tesouraria no mercado.",
        "Controladores — formulário 'Posição Consolidada', grupo Controlador.",
        "Administradores — formulário 'Posição Consolidada', grupos Conselho de Administração "
        "+ Diretoria (Lei 6.404). Conselho Fiscal e Órgãos Técnicos ficam em Detalhe (Aba = Outros).",
        "Detalhe — uma linha por operação lida de cada documento vigente, inclusive as que "
        "não entram nos resumos.",
        "Documentos — todos os documentos encontrados, reapresentações substituídas, "
        "documentos ainda não entregues e a conferência de saldo de cada um.",
        "",
        "REGRAS",
        "1. Só AÇÕES negociadas no mercado (compra/venda à vista ou a termo) entram nos "
        "resumos. Vesting/ILP, liquidação de swap, empréstimo de ações, subscrição, "
        "bonificação, posse/saída de cargo, debêntures, opções e ADRs ficam só em Detalhe.",
        "2. Compra/venda declarada sem preço vira 'Compra s/ preço' / 'Venda s/ preço': aparece "
        "em Detalhe e fecha a conferência de saldo, mas fica fora de Qtd, Valor e Preço médio. "
        "Par 'compra a termo' + 'venda à vista' no mesmo dia e quantidade vira 'Rolagem a termo' "
        "(a posição não muda) e também fica fora dos resumos.",
        "3. Posições em ações da CONTROLADORA (ex.: administradores da VIVT3 com ações da "
        "Telefónica S.A.) não são ações da companhia e ficam em Aba = Outros.",
        "4. Reapresentação: por empresa, mês e formulário vale só o documento mais recente.",
        "5. Qtd movimentada e Valor total são fórmulas SUMIFS sobre a aba Detalhe; "
        "Preço médio = Valor total / Qtd (ponderado).",
        "6. Conferência de saldo: saldo inicial + compras - vendas deve igualar o saldo "
        "final. 'OK*' = diferença coberta por operações não-mercado do mesmo grupo; "
        "'VERIFICAR' (amarelo) = conferir no PDF.",
        "",
        "ATUALIZAR: py buybacks_cvm.py (ver README.md, seção 'Planilha de recompras e insiders').",
    ]
    for i, txt in enumerate(linhas, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = Font(name="Arial", size=12 if i == 1 else 10, bold=(i == 1 or txt in (
            "ABAS", "REGRAS")))
        c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


def print_summary(ops_rows: List[Dict]) -> None:
    if not ops_rows:
        return
    ultimo = max(r["meta"]["ref"] for r in ops_rows)
    log(f"Resumo {ultimo[5:]}/{ultimo[:4]} (mercado, ações):")
    for aba in (ABA_EMPRESA, ABA_CONTROLADORES, ABA_ADMINISTRADORES):
        tot: Dict[Tuple[str, str], List[float]] = {}
        for r in ops_rows:
            if (r["meta"]["ref"] == ultimo and r["aba"] == aba and r["vm"] == "Ações"
                    and r["op"] in (OP_COMPRA, OP_VENDA)):
                t = tot.setdefault((r["meta"]["ticker"], r["op"]), [0, 0.0])
                t[0] += r["qtd"]
                t[1] += r["volume"]
        itens = [f"{tk} {op.lower()} {q:,.0f} ações R$ {v / 1e6:,.1f} mi"
                 for (tk, op), (q, v) in sorted(tot.items())]
        log(f"  {aba}: " + ("; ".join(itens) if itens else "nada"))


# ============================================================================
# MAIN
# ============================================================================

def load_cache() -> Dict[str, Dict]:
    if not CACHE_FILE.exists():
        return {}
    with open(CACHE_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def save_cache(cache: Dict[str, Dict]) -> None:
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=1, sort_keys=True)


def refresh(reparse: bool = False) -> Dict[str, Dict]:
    session = cvm.make_session()
    docs = list_documents(session)
    cache = load_cache()

    for prot, doc in sorted(docs.items(), key=lambda kv: kv[1]["entrega"]):
        entry = cache.get(prot)
        parsed_ok = (entry and (entry.get("parse") or {}).get("parser") == PARSER_VERSION
                     and not entry["parse"].get("erro"))  # erro de download é retentado
        if parsed_ok and not reparse:
            entry["meta"] = doc
            continue
        try:
            parsed = parse_pdf(fetch_pdf(session, doc))
            log(f"Lido {doc['ticker']} {doc['ref']} {doc['formulario']} ({prot}): "
                f"{len(parsed['secoes'])} seções, "
                f"{sum(len(s['ops']) for s in parsed['secoes'])} operações")
        except Exception as e:
            log(f"ERROR: {doc['ticker']} {doc['ref']} ({prot}): {e}")
            parsed = {"parser": PARSER_VERSION, "erro": str(e), "secoes": []}
        cache[prot] = {"meta": doc, "parse": parsed}

    save_cache(cache)
    return cache


def dump_ops(cache: Dict[str, Dict]) -> None:
    seen: Dict[Tuple[str, str], int] = {}
    for entry in cache.values():
        for sec in entry["parse"].get("secoes", []):
            for op in sec["ops"]:
                key = (valor_mobiliario(op["desc"]), classificar_operacao(op), op["desc"])
                seen[key] = seen.get(key, 0) + 1
    for (vm, cls, desc), n in sorted(seen.items()):
        print(f"{vm:12} {cls:7} {n:4}  {desc}")


def main() -> int:
    for _stream in (sys.stdout, sys.stderr):
        try:
            _stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--reparse", action="store_true", help="relê todos os PDFs")
    ap.add_argument("--offline", action="store_true",
                    help="não consulta CVM nem B3; só regera xlsx e painel a partir do cache")
    ap.add_argument("--dump-ops", action="store_true",
                    help="lista descrições de operação distintas e sua classificação")
    ap.add_argument("--sem-painel", action="store_true",
                    help="não regera o painel HTML (buybacks_dashboard.html)")
    args = ap.parse_args()

    cache = load_cache() if args.offline else refresh(reparse=args.reparse)
    if args.dump_ops:
        dump_ops(cache)
        return 0

    ops_rows, doc_rows = build_rows(cache)
    status = 0
    try:
        write_xlsx(ops_rows, doc_rows, OUTPUT_XLSX)
    except PermissionError:
        # No Windows, o Excel trava o arquivo aberto. O painel sai mesmo assim.
        log(f"ERROR: {OUTPUT_XLSX.name} está aberto (Excel?) — feche e rode de novo "
            "com --offline para regravar a planilha.")
        status = 1
    verificar = [d for d in doc_rows
                 if d["usado"] and not d["conferencia"].startswith("OK")]
    log(f"Planilha: {OUTPUT_XLSX.name} ({len(ops_rows)} operações, "
        f"{len(doc_rows)} documentos, {len(verificar)} para verificar).")
    for d in verificar:
        m = d["meta"]
        log(f"  VERIFICAR {m['ticker']} {m['ref']} {m['formulario']} ({m['protocolo']}): "
            f"{d['conferencia']}")
    print_summary(ops_rows)

    if not args.sem_painel:
        # O painel usa o mesmo cache + cotação da B3 e ITR/DFP (buybacks_mercado).
        # Falha nele não pode derrubar a planilha, que já foi gravada.
        try:
            import build_buybacks_dashboard
            build_buybacks_dashboard.build(offline=args.offline)
        except Exception as e:
            log(f"ERROR: painel não foi regerado: {e}")
            return 1
    return status


if __name__ == "__main__":
    raise SystemExit(main())
